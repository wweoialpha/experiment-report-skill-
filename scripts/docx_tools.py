"""Small reusable helpers for experiment reports; run with --help."""
import argparse
from copy import deepcopy
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile


def compile_formulas(formulas, pandoc='pandoc'):
    """Return native OMML objects in input order; no document is modified."""
    from docx import Document
    if not isinstance(formulas, list) or not formulas or any(
        not isinstance(s, str) or not s.strip() or '$$' in s for s in formulas
    ):
        raise ValueError('Expected a nonempty JSON list of LaTeX strings without $$ delimiters')
    executable = shutil.which(pandoc)
    if not executable:
        raise FileNotFoundError('Pandoc not found; pass --pandoc with its executable path')
    with tempfile.TemporaryDirectory(prefix='report_math_') as tmp:
        source, target = Path(tmp) / 'math.md', Path(tmp) / 'math.docx'
        source.write_text('\n\n'.join('$$\n' + s + '\n$$' for s in formulas), encoding='utf-8')
        proc = subprocess.run([executable, str(source), '-o', str(target)],
                              capture_output=True, text=True, encoding='utf-8',
                              errors='replace', timeout=60)
        if proc.returncode or proc.stderr.strip():
            raise RuntimeError('Formula conversion requires review: ' + proc.stderr)
        objects = Document(target)._element.xpath('//m:oMath')
        if len(objects) != len(formulas):
            raise RuntimeError('Formula count mismatch; conversion is incomplete')
        return [deepcopy(obj) for obj in objects]


def formula_bank(args):
    from docx import Document
    from docx.oxml import OxmlElement
    destination = Path(args.output)
    if destination.exists():
        raise FileExistsError(f'Refusing to overwrite {destination}')
    objects = compile_formulas(json.loads(Path(args.input).read_text(encoding='utf-8-sig')), args.pandoc)
    doc = Document()
    for obj in objects:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1
        math_paragraph = OxmlElement('m:oMathPara')
        math_paragraph.append(obj)
        paragraph._p.append(math_paragraph)
    doc.save(destination)
    print(f'Created {destination}: {len(objects)} editable formulas')


def inherited(run, attribute):
    value = getattr(run.font, attribute)
    if value is not None:
        return value
    for style in (run.style, run._parent.style):
        seen = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            value = getattr(style.font, attribute)
            if value is not None:
                return value
            style = style.base_style
    return None


def check(args):
    """Read-only checks for the skill's current font-size and geometry defaults."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    doc = Document(args.input)
    findings = []
    title = next((p._p for p in doc.paragraphs if p.text.strip()), None)
    for index, node in enumerate(doc._element.xpath('//w:body//w:p'), 1):
        p = Paragraph(node, doc._body)
        expected = 22 if node is title else 14 if re.fullmatch(r'\[.+\]', p.text.strip()) else 12
        for run in p.runs:
            if not run.text.strip():
                continue
            size = inherited(run, 'size')
            if size is None or abs(size.pt - expected) > 0.1:
                findings.append(f'paragraph {index}: font size {size.pt if size else "unresolved"}, expected {expected}')
                break
        if node is title and node.xpath('.//w:br'):
            findings.append('title: explicit line/page break')
        if re.search(r'\[\[|\]\]|\$\$|\\(?:frac|sqrt|lambda|begin)\b', p.text):
            findings.append(f'paragraph {index}: possible unconverted math')
    widths = [s.page_width - s.left_margin - s.right_margin for s in doc.sections]
    for index, table in enumerate(doc.tables, 1):
        grid = sum(int(c.get(qn('w:w'))) for c in table._tbl.tblGrid)
        if grid * 635 > max(widths) + 1270:
            findings.append(f'table {index}: grid wider than every section text area')
        for row in table.rows:
            total = sum(int(c.tcPr.tcW.w) for c in row._tr.tc_lst if c.tcPr is not None and c.tcPr.tcW is not None)
            if abs(total - grid) > 3:
                findings.append(f'table {index}: cell and grid widths differ')
                break
    for index, shape in enumerate(doc.inline_shapes, 1):
        if shape.width > max(widths):
            findings.append(f'image {index}: wider than every section text area')
    print(json.dumps({'file': str(Path(args.input)), 'math_objects': len(doc._element.xpath('//m:oMath')),
                      'tables': len(doc.tables), 'images': len(doc.inline_shapes), 'findings': findings,
                      'scope': 'Body text sizes, explicit title breaks, residual LaTeX, basic widths. '
                               'Not a complete style, numerical, or visual check; font faces, complex inheritance, '
                               'section-local placement, pagination and formula appearance require inspection.'},
                     ensure_ascii=False, indent=2))
    return 1 if findings else 0


def render(args):
    """Reuse the installed documents renderer; fix Windows file URI construction."""
    source = Path(args.renderer).resolve()
    if not source.is_file() or source == Path(__file__).resolve():
        raise FileNotFoundError('Pass the installed documents skill render_docx.py via --renderer')
    destination = Path(args.output)
    if destination.exists() and any(destination.iterdir()):
        raise FileExistsError('Render output directory must be new or empty')
    spec = importlib.util.spec_from_file_location('report_renderer', source)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    if not callable(getattr(module, '_run_cmd', None)) or not callable(getattr(module, 'main', None)):
        raise RuntimeError('Renderer interface changed; inspect the installed renderer before using this helper')
    def run(cmd, env, verbose):
        fixed = []
        for argument in cmd:
            prefix = '-env:UserInstallation=file://'
            if os.name == 'nt' and argument.startswith(prefix):
                raw = argument[len(prefix):]
                if re.match(r'^[A-Za-z]:[\\/]', raw):
                    argument = '-env:UserInstallation=' + Path(raw).as_uri()
            fixed.append(argument)
        if verbose:
            print(fixed)
        return subprocess.run(fixed, env=env, capture_output=True, text=True,
                              encoding='utf-8', errors='replace', timeout=60)
    module._run_cmd = run
    previous = sys.argv
    sys.argv = [str(source), str(Path(args.input).resolve()), '--output_dir', str(destination.resolve())]
    if args.emit_pdf:
        sys.argv.append('--emit_pdf')
    try:
        module.main()
    finally:
        sys.argv = previous


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    f = sub.add_parser('formulas', help='Compile a JSON LaTeX list into an editable DOCX formula bank')
    f.add_argument('input'); f.add_argument('output'); f.add_argument('--pandoc', default='pandoc')
    f.set_defaults(action=formula_bank)
    c = sub.add_parser('check', help='Read-only partial format checks; JSON to stdout')
    c.add_argument('input'); c.set_defaults(action=check)
    r = sub.add_parser('render', help='Render using the installed documents skill')
    r.add_argument('input'); r.add_argument('output'); r.add_argument('--renderer', required=True)
    r.add_argument('--emit-pdf', action='store_true'); r.set_defaults(action=render)
    args = parser.parse_args()
    try:
        return args.action(args) or 0
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError) as exc:
        parser.exit(2, f'{type(exc).__name__}: {exc}\n')


if __name__ == '__main__':
    sys.exit(main())
